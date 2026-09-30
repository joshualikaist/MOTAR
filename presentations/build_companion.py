#!/usr/bin/env python3
"""Build the editable paper DOCX and 12-slide overview; no simulator or experiment runs.

Install requirements-authoring.txt in a separate environment, then run this script.
--check regenerates in memory and compares both Office files and their source manifest.
Scientific values come from the source-bound final figure manifest, whose upstream
SHA-256 values are verified before authoring. Never edit an experiment to update prose.
"""
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import argparse
import json
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
STAMP = "2026-09-26"
PAPER = ROOT / "paper" / ("MOTAR_draft_" + STAMP + ".md")
FIGURES = ROOT / "docs/assets/paper/final"
DECK = "presentations/MOTAR_overview_" + STAMP + ".pptx"
DOCX = "paper/MOTAR_draft_" + STAMP + ".docx"
MANIFEST = "presentations/companion_manifest.json"
INK, BLUE, MUTED, RULE, BG = "1F1F1F", "2F6599", "626A73", "D7DCE0", "FBFAF7"
FONT = "Noto Sans CJK KR"


def pp(value, digits=2):
    return ("{:+." + str(digits) + "f}").format(value).replace("-", "−")


def interval(bounds, digits=2):
    return "[" + ", ".join(pp(x, digits) for x in bounds) + "]"


def evidence():
    manifest = json.loads((FIGURES / "manifest.json").read_text())
    for path, digest in manifest["sources"].items():
        if sha256((ROOT / path).read_bytes()).hexdigest() != digest:
            raise ValueError("Evidence source changed: " + path)
    return manifest["values"]


def deterministic_office(data):
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as src, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info = zipfile.ZipInfo(name, (2026, 9, 26, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dst.writestr(info, src.read(name))
    return output.getvalue()


def metadata(core, title):
    core.title = title
    core.author = "Minseok Lee; Hyondong Oh"
    core.subject = "MOTAR simulation evidence; independent seven-seed target-motion replication"
    core.keywords = "MOTAR; UAV; tracking; close approach; simulation; replication"
    core.comments = "Review draft. Evidence closed. Generated from repository sources."
    core.created = core.modified = datetime(2026, 9, 26, tzinfo=timezone.utc)
    core.last_modified_by = "MOTAR companion builder"
    core.revision = 1


def build_pptx(v):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as Shape, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    metadata(prs.core_properties, "MOTAR — Independent Target-Motion Replication")
    tm, p10, safe, d8 = v["target_motion"], v["p10"], v["safety"], v["d8b"]
    contrasts = tm["contrasts_vs_H"]
    slide_records = []

    def text(slide, x, y, w, h, content, size=20, color=INK, bold=False, align=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for index, line in enumerate(content.split("\n")):
            p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            p.text = line
            p.space_after = Pt(5)
            if align is not None:
                p.alignment = align
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = RGBColor.from_string(color)
                east = OxmlElement("a:ea")
                east.set("typeface", FONT)
                run._r.get_or_add_rPr().append(east)
        return box

    def line(s, x1, y1, x2, y2, color=RULE, width=1):
        shp = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shp.line.color.rgb = RGBColor.from_string(color)
        shp.line.width = Pt(width)
        return shp

    def dot(s, x, y, radius=.055, color=BLUE, fill=True):
        shp = s.shapes.add_shape(Shape.OVAL, Inches(x-radius), Inches(y-radius), Inches(radius*2), Inches(radius*2))
        shp.line.color.rgb = RGBColor.from_string(color)
        shp.line.width = Pt(1)
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(color if fill else BG)
        return shp

    def rect(s, x, y, w, h, fill=BG):
        shp = s.shapes.add_shape(Shape.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        shp.line.color.rgb = RGBColor.from_string(RULE)
        shp._element.spPr.append(OxmlElement("a:effectLst"))
        return shp

    def new(kicker, title, key, notes, sources):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        n = len(prs.slides)
        text(slide, .62, .32, 11.8, .3, kicker.upper(), 11, BLUE, True)
        text(slide, .62, .92, 12.05, 1.1, title, 29, INK, True)
        line(slide, .62, 6.53, 12.7, 6.53)
        text(slide, .62, 6.69, 12, .36, key, 14, BLUE, True)
        text(slide, .62, 7.14, 11.7, .2, "MOTAR  /  Research draft  /  26 September 2026  /  Simulation only", 9, MUTED)
        text(slide, 12.05, 7.1, .65, .25, f"{n:02d} / 12", 9, MUTED, align=PP_ALIGN.RIGHT)
        note = notes + "\n\nSources (repository-relative):\n" + "\n".join(sources)
        slide.notes_slide.notes_text_frame.text = note
        slide_records.append({"slide": n, "title": title, "key": key, "sources": sources})
        return slide

    sources_tm = ["results/target_motion_replication_7seed/analysis.json", "results/target_motion_replication_7seed/replication_record.json"]
    sources_p10 = ["results/perception_p10_seed_replication_2026-09-10/summary.json", "docs/plans/perception_p9_p10_execution_2026-09-09.md"]
    sources_safety = ["results/independent_verification_2026-09-07/recomputed.json", "results/independent_verification_2026-09-07/PAPER_METHOD_EXPERIMENT.md"]
    sources_d8 = ["results/dynamic_mesh_policy_sensitivity_d8b_2026-09-13/summary.json", "docs/d8c_perception_path_audit_2026-09-24.md"]

    s = new("MOTAR / Moving Object Tracking And Rendezvous", "표적 운동의 효과를 독립 n=7 실험으로 재현했다", "Frozen policy F  ·  7 fresh seeds  ·  Simulation-only evidence",
            f"오늘 발표의 중심은 새로운 학습법이 아니라 고정 정책 F의 평가 계약입니다. 표적 운동 효과는 환경별 할당량(per-environment quota)을 교정한 독립 실험에서 재현됐습니다. 화면의 세 숫자는 역사적 학습 표적 H 대비 capture 차이이며, 새 시드 {len(tm['seeds'])}개, 112셀, {tm['episodes']:,}개 primary episode에서 나온 값입니다. 원본 n=3 캠페인은 별도 역사적 근거로 보존하고 합산하지 않습니다. 모든 정책 평가는 시뮬레이션이며 실비행 주장은 없습니다. 전체 실험 프로그램은 종료됐고, 이 발표는 새 실험 계획이 아니라 기록을 정리한 것입니다.", sources_tm)
    text(s, .68, 2.05, 11.8, .85, "Independent replication of target-motion effects\nin UAV close-approach simulation", 23, MUTED)
    for i, arm in enumerate(("E0", "E1", "E2")):
        x = .72 + i*4.05
        text(s, x, 3.55, 3.6, .4, {"E0":"Static − historical", "E1":"Constant velocity − historical", "E2":"Obstacle-aware − historical"}[arm], 16, MUTED)
        text(s, x, 4.05, 3.5, .85, pp(contrasts[arm]["mean_pp"])+" pp", 41, BLUE, True)
    text(s, .72, 5.68, 11.6, .5, "이민석 · 오현동  |  Minseok Lee · Hyondong Oh\nKAIST · Department of Mechanical Engineering", 15)

    s = new("01 / Research question", "정책을 고정하면, 어떤 가정의 변화가 성능을 바꾸는가?", "Research question  ·  Contract sensitivity, not a causal ranking",
            "MOTAR는 지각 불확실성, 정책 민감도, 안전 필터, 표적 운동 분포, 관측 계약을 연결된 연구 질문으로 다룹니다. 정책 가중치는 평가 동안 고정하고, 한 번에 하나의 계약만 바꿔 capture, crash, timeout의 변화를 봅니다. 네 비교군은 서로 다른 실험 모듈이며 하나의 인과 사슬이 아닙니다. 지각, 안전, 표적 운동은 policy F(ep25000 + riskcap), 관측 계약은 policy R(ref5in D1 ep1900)이라는 별도 체크포인트를 사용하므로 효과 크기를 서로 순위화하지 않습니다.", ["ARCHITECTURE.md", "docs/EVIDENCE.md"])
    rect(s, 4.75, 3.12, 3.7, 1.35, "FFFFFF")
    text(s, 5, 3.38, 3.2, .75, "FROZEN POLICY\nmeasured response", 23, BLUE, True, PP_ALIGN.CENTER)
    for x,y,label,ex,ey in [(1,2.38,"Perception error",4.75,3.4),(9,2.38,"Safety filter",8.45,3.4),(1,4.75,"Target motion",4.75,4.1),(9,4.75,"Observation contract",8.45,4.1)]:
        text(s,x,y,3.25,.8,label,21)
        line(s,x+1.55,y+.45,ex,ey)

    s = new("02 / System overview", "센서 관측에서 제어까지의 경계를 명시한다", "Actor observation  ·  Ground-truth target state is excluded",
            "카메라 기반 검출과 짧은 track 이력, 72×4 LiDAR 장애물 관측을 17개 토큰으로 묶어 Transformer PPO 정책에 전달합니다. 정책이 낸 속도 명령은 LiDAR 속도 필터를 거쳐 Lee 속도 제어기로 가며, RL 한 스텝은 0.1 s입니다. 표적의 GT 위치, 방위, 거리, 가시성은 감독, 보상, 종료, critic, 평가에만 허용하며 actor 관측에는 넣지 않습니다. 실제 영상 분석은 오프라인 측정 분기이고 실비행 폐루프 검증은 없습니다.", ["ARCHITECTURE.md", "docs/assets/paper/final/fig1-motar-overview.svg"])
    labels = [("Sensors","Camera + LiDAR"),("Perception","Detection + history"),("Policy","Transformer PPO"),("Filter","LiDAR speed cap"),("Control","Lee velocity")]
    for i,(title,sub) in enumerate(labels):
        x=.7+i*2.51
        rect(s,x,3.05,2.18,1.55,"FFFFFF")
        text(s,x+.15,3.32,1.88,.4,title,22,BLUE,True,PP_ALIGN.CENTER)
        text(s,x+.1,3.99,1.98,.4,sub,13,MUTED,False,PP_ALIGN.CENTER)
        if i<4: text(s,x+2.2,3.52,.3,.4,"→",20,BLUE)
    line(s,1.8,4.8,11.9,4.8)
    text(s,3.12,5.22,7.1,.6,"Simulated obstacle field → next observation",21,MUTED,False,PP_ALIGN.CENTER)

    s = new("03 / Evidence chain", "각 결과의 정책과 추론 단위를 함께 읽는다", "Separate evidence modules  ·  Policy F ≠ policy R",
            "이 표의 순서는 인과 강도의 순위가 아닙니다. 각 행은 서로 다른 실험 모듈이고 추론 단위가 다릅니다. E3S 거리 proxy는 한 비행의 아홉 블록이고, P9 비용은 두 평가 시드에서 모은 episode-level 구간입니다. P10 재적응은 세 학습 시드, 안전 필터 비교는 세 평가 시드가 공유되는 15셀입니다. 표적 운동은 일곱 평가 시드가 추론 단위이며, episode 수를 독립 표본 수로 읽지 않습니다. D8b는 다른 정책 R을 사용하므로 손실 크기를 F의 효과와 직접 순위화할 수 없습니다.", ["docs/assets/paper/final/manifest.json"])
    rows = [("Range proxy",f"{v['e3s']['median_of_block_medians_pct']:.1f}%","1 flight / 9 blocks"),("Pixel / temporal cost",pp(p10['frozen_cost_pp'])+" pp","F / 2 evaluation seeds*"),("Readaptation",pp(p10['readapt_mean_pp'])+" pp","3 training seeds / INCONCLUSIVE"),("Configured filter",pp(safe['pooled_pp'],4)+" pp crash","F / 3 evaluation seeds"),("Target-motion replication","n = 7","F / paired evaluation seeds"),("Observation contract",pp(d8['primary_pp'],3)+" pp","R / 3 paired evaluation seeds")]
    for i,(label,value,unit) in enumerate(rows):
        y=2.11+i*.62
        text(s,.8,y,3.75,.4,label,18)
        text(s,4.8,y,2.85,.4,value,20,BLUE,True)
        text(s,8,y,4.6,.4,unit,16,MUTED)
        line(s,.8,y+.47,12.5,y+.47,width=.6)
    text(s,.8,6.04,11.5,.3,"* P9 interval is episode-level; the repeated frozen arms are a determinism check.",12,MUTED)

    s = new("04 / Perception measurement → policy sensitivity", "거리 오차 측정과 정책 오차 주입은 별도 분기다", "P10 readaptation  ·  INCONCLUSIVE; no net benefit established",
            "왼쪽 ETH ds5의 6.2%는 한 비행 3,107프레임을 아홉 블록으로 집계한 거리 proxy입니다. 오른쪽 NPS P8은 1,310개의 exactly-one-GT 검증 프레임 27구간에서 픽셀 및 시간 오차를 측정합니다. P9는 HIT/FALSE_LOCK/NO_LOCK 전이, 중심점 오프셋, decode+inference 지연을 주입하며 analytic surface range/confidence를 유지합니다. ETH 거리 오차는 주입하지 않습니다. P9의 −4.57 구간은 두 평가 시드에서 episode-level이고, P10 +0.73 구간은 세 학습 시드 기반입니다. 양의 평균만으로 재학습 이득을 확정하지 않습니다.", ["results/eth_ds5_e3s_2026-09-10/run/e3s_result.json"]+sources_p10)
    text(s,.8,2.25,4.05,.45,"ETH ds5 / range-proxy measurement",17,MUTED)
    text(s,.8,3.12,4,.85,f"{v['e3s']['median_of_block_medians_pct']:.1f}%",44,BLUE,True)
    text(s,.8,4.13,3.7,1.3,"Median of 9 block medians\nOne flight; correlated frames\nNot the P9 injected statistic",18)
    line(s,5,2.25,5,5.95)
    text(s,5.52,2.25,7.05,.5,"NPS P8 → pixel / temporal injection → policy F",19,BLUE,True)
    for y,name,val,ci,unit in [(3.25,"Frozen-policy cost",p10['frozen_cost_pp'],p10['frozen_cost_ci'],"episode-level 95% CI"),(4.77,"Readaptation",p10['readapt_mean_pp'],p10['readapt_ci'],"training-seed 95% CI")]:
        text(s,5.52,y,3.3,.4,name,19)
        text(s,9,y-.1,3.5,.6,pp(val)+" pp",30,BLUE,True)
        text(s,5.52,y+.65,6.9,.5,interval(ci)+" pp  ·  "+unit,15,MUTED)

    s = new("05 / Safety geometry", "구현된 arc 필터는 riskcap보다 crash를 줄였다", "Configured comparison  ·  Geometry and stopping law both change",
            "헤드라인 −1.4903 pp는 arc-clearance−riskcap의 구성된 필터 비교입니다. 기하만의 인과 효과가 아닙니다. 점 15개는 세 평가 시드 × 다섯 밀도이므로 독립 반복 15회로 부르면 안 됩니다. 등록된 cell-pooled CI는 [−1.8981,−1.0826], seed-level t CI는 [−2.42,−0.57] pp입니다. arc−stopcap은 정지 법칙을 고정한 별도 비교입니다. 속도 필터는 안전 증명이나 DWA planner가 아니며 capture와 timeout의 교환관계도 봐야 합니다.", sources_safety)
    text(s,.78,2.3,5.5,.45,"Arc-clearance − riskcap / crash",21)
    text(s,.78,3.03,5.5,.95,pp(safe['pooled_pp'],4)+" pp",43,BLUE,True)
    text(s,.78,4.18,5.3,1.4,"15 / 15 cells lower\n3 evaluation seeds\nNo collision guarantee",21)
    x0,x1,y0,y1=7,12.3,2.65,5.52
    def sx(value): return x0+(value+3.5)/3.7*(x1-x0)
    for tick in (-3,-2,-1,0):
        line(s,sx(tick),y0-.1,sx(tick),y1)
        text(s,sx(tick)-.3,y1+.16,.6,.35,str(tick).replace("-","−"),14,MUTED,align=PP_ALIGN.CENTER)
    for i,seed in enumerate(safe['seeds']):
        y=2.98+i*.91
        text(s,5.95,y-.18,.95,.3,"seed "+str(seed),11,MUTED)
        cells=[r for r in safe['cells'] if r['seed']==seed]
        for j,row in enumerate(cells): dot(s,sx(row['delta_pp']),y+(j-2)*.065,.057,fill=False)
    text(s,6.4,5.97,6.3,.35,"Crash difference per cell (pp)  ·  seed-level 95% CI "+interval(safe['seed_ci']),13,MUTED)

    s = new("06 / Central result · Independent replication", "교정된 할당량에서도 세 방향성 대비가 재현됐다", "E2 positive: 5/7 at 160 bars, 4/7 at 205 bars  ·  Density-specific analyses are exploratory",
            f"중심 결과입니다. 새 시드 4104–4110, 네 arm과 네 밀도의 112셀에서 환경별 처음 완료된 16개 에피소드 ×128환경을 선택해 셀당 2,048개, 총 229,376개가 primary입니다. 각 점은 시드별 네 밀도 평균 대비이고 굵은 구간은 BCa95입니다. " + ", ".join(f"{a} {tm['arm_mean_capture_pct'][a]:.2f}%" for a in ("H","E0","E1","E2")) + ". exact p=0.015625, Holm p=0.046875로 일곱 시드 부호반전 검정의 해상도 경계입니다. highly significant라고 하지 않습니다. E2는 160 bars 5/7, 205 bars 4/7만 양의 방향이며 밀도별 분석은 exploratory입니다. 원본 n=3과 합치지 않습니다.", sources_tm)
    text(s,.75,2.12,11.9,.45,f"7 fresh seeds  ·  112 cells  ·  {tm['episodes']:,} primary episodes",20,MUTED)
    xa,xb=3.47,8.7
    def tx(value): return xa+(value+7)/10*(xb-xa)
    for tick in (-6,-4,-2,0,2):
        line(s,tx(tick),2.88,tx(tick),5.51,INK if tick==0 else RULE,width=.8)
        text(s,tx(tick)-.3,5.57,.6,.3,str(tick).replace("-","−"),13,MUTED,align=PP_ALIGN.CENTER)
    for i,(arm,label) in enumerate((("E0","Static"),("E1","Constant velocity"),("E2","Obstacle-aware"))):
        c=contrasts[arm]; y=3.21+i*.87
        text(s,.8,y-.18,2.6,.5,arm+"  "+label,18)
        for j,value in enumerate(c['seed_pp']): dot(s,tx(value),y-.16+(j%3)*.15,.05,MUTED,False)
        lo,hi=c['bca95_pp']; line(s,tx(lo),y,tx(hi),y,BLUE,3)
        dot(s,tx(c['mean_pp']),y,.079,BLUE)
        text(s,9.02,y-.3,3.6,.48,pp(c['mean_pp'])+" pp",28,BLUE,True)
        text(s,9.02,y+.21,3.6,.3,interval(c['bca95_pp'])+"  BCa95",14,MUTED)
    text(s,.78,6.01,11.9,.3,"All 7 / 7 directions agree  ·  exact p = 0.015625  ·  Holm p = 0.046875 (resolution boundary)",14,MUTED)

    s = new("07 / Observation-contract sensitivity", "관측 계약 변화에서 손실을 확인했고, 원인은 미검증이다", "MATERIAL_LOSS  ·  causality_vs_d8b = NOT_TESTED  ·  Policy R, not F",
            "D8b는 같은 R 정책과 검출기, task와 paired seeds에서 관측 treatment를 바꿉니다. mesh-shaded−analytic은 −48.967 pp, 95% seed-t CI [−50.113,−47.821]입니다. 특정 shading, geometry, detector 또는 policy 결함을 독립 원인으로 입증하지 않았습니다. 다른 정책 F의 효과와 크기를 단순 순위화하지 않습니다. 관측 계약 audit을 재학습보다 먼저 검토할 근거이며 PPO가 필요하다는 결론이 아닙니다. D8c는 읽기 전용 감사이고 새 실험은 없습니다.", sources_d8)
    text(s,.8,2.32,5.5,.65,"Same detector + frozen policy R\nDifferent observation treatment",22)
    for i,label in enumerate(("Analytic proxy","Mesh flat","Mesh shaded")):
        rect(s,.8+i*1.88,3.72,1.68,1.15,"FFFFFF")
        text(s,.92+i*1.88,3.97,1.43,.7,label,17,BLUE,True,PP_ALIGN.CENTER)
    line(s,6.64,2.4,6.64,5.75)
    text(s,7.2,2.78,5.3,.5,"Mesh shaded − analytic / capture",18,MUTED)
    text(s,7.2,3.57,5.3,.85,pp(d8['primary_pp'],3)+" pp",40,BLUE,True)
    text(s,7.2,4.75,5.4,.8,interval(d8['primary_ci'],3)+" pp\n95% CI · 3 paired evaluation seeds",17)

    s = new("08 / Relation to published systems", "문헌은 평가 계약을 비교하는 맥락으로 사용한다", "Class A matched external comparisons = 0  ·  No performance ranking",
            "NavRL/NavRL++은 navigation과 sim-to-real, Fast/Elastic/YOPOv2-Tracker는 추적 planning 또는 perception-to-action, OPEN은 multi-UAV pursuit-evasion의 맥락입니다. 각 논문의 정보 입력, 플랫폼, target assumptions, denominator가 다릅니다. 외부 성능 수치를 빼거나 우월성을 주장하지 않습니다. Elastic Tracker B0는 환경 문제로 reproduction 실패이며 benchmark가 아닙니다. 자세한 primary references는 논문 R1–R6 및 evidence ledger에 있습니다.", ["docs/quantitative_positioning_registry.json", "paper/MOTAR_draft_2026-09-26.md"])
    for i,(name,scope) in enumerate([("NavRL / NavRL++","Goal navigation; different task contract"),("Fast / Elastic / YOPOv2-Tracker","Tracking; different observations and outcomes"),("OPEN","Multiple pursuers; privileged state assumptions"),("MOTAR","Frozen-policy sensitivity; independent replication")]):
        y=2.35+i*.82
        text(s,.8,y,5.2,.45,name,20,BLUE if i==3 else INK,i==3)
        text(s,6.12,y,6.45,.55,scope,18,MUTED)
        line(s,.8,y+.63,12.45,y+.63)
    text(s,.8,6.02,11.8,.3,"Elastic Tracker: B0 environmental failure; not benchmarked.",15,MUTED)

    s = new("09 / Limitations", "관측된 효과의 범위를 결론 안에 남긴다", "No real flight  ·  No matched external benchmark  ·  No validated min-distance result",
            "simulation only이며 실비행과 외부 Class A benchmark는 없습니다. E2의 방향은 고밀도에서 일정하지 않고 n=7 검정은 해상도 경계입니다. D8b 원인은 미검증이고 P10은 inconclusive입니다. 최소거리는 step-end sampling과 swept capture의 의미 불일치로 WITHHELD_SEMANTIC_MISMATCH 상태입니다. raw는 normal Git 밖에 보관하며 공개 archive download는 없습니다. manifest와 hash가 보존되어도 raw 공개와 같은 뜻이 아닙니다.", sources_tm+sources_p10+sources_d8)
    pairs=[("Scope","Simulation only / no real flight"),("Inference","7-seed resolution boundary / high-density E2 limit"),("Mechanism","D8b NOT_TESTED / P10 INCONCLUSIVE"),("Measurement","Min-distance WITHHELD_SEMANTIC_MISMATCH"),("Access","Raw evidence outside Git / no public archive download")]
    for i,(name,limit) in enumerate(pairs):
        y=2.3+i*.74
        text(s,.8,y,2.55,.45,name,20,BLUE,True)
        text(s,3.45,y,9,.5,limit,19)
        line(s,.8,y+.58,12.45,y+.58)

    s = new("10 / What is ready for paper", "원본 결과와 교정 과정을 보존한 논문 패키지를 만들었다", "Current evidence: n = 7  ·  Historical n = 3 remains separate and unpooled",
            "원본 9월18일 n=3, stopping-window bias 감사와 canary, 9월25–26일 독립 n=7을 구분합니다. Canary는 instrumentation validation으로 과학적 replication이 아닙니다. 같은 n=7에서 quota와 legacy estimator를 비교할 수 있으나 시드가 다른 campaign 전체 차이를 estimator만으로 설명할 수는 없습니다. 최신 문서와 그림, PPT, 공개 CPU 검증은 편집 패키지입니다.", sources_tm+["paper/PAPER_TODO_2026-09-26.md"])
    for i,(date,title,scope) in enumerate([("18 SEP","Original / n = 3","Legacy pooled window\nPreserved historical evidence"),("AUDIT","Bias diagnosis + canary","Instrumentation validation\nNot scientific replication"),("25–26 SEP","Replication / n = 7","Per-environment quota\nCurrent confirmatory evidence")]):
        x=.8+i*4.16
        dot(s,x+.12,3.02,.085,BLUE)
        if i<2: line(s,x+.25,3.02,x+4.15,3.02,BLUE)
        text(s,x,2.35,3.9,.4,date,15,BLUE,True)
        text(s,x,3.47,3.82,.7,title,23,INK,True)
        text(s,x,4.5,3.85,1.1,scope,18,MUTED)
    text(s,.8,5.83,11.8,.4,"Paper draft + 6 figures + 12-slide deck + source-bound public checks",20)

    s = new("11 / Next steps · Editorial work only", "남은 작업은 문헌·서식·전체 diff 검토다", "Experimental programme closed  ·  No new training, seeds or GPU evaluation",
            "다음 단계는 citation metadata, 저자와 소속, venue template, figure print export, tables와 data availability statement의 검토입니다. 새 PPO, extra seeds, D8c experiment, Elastic Tracker run은 시작하지 않습니다. 전체 diff와 검증 결과를 제시하고 커밋 전 검토를 기다립니다. 이 단계의 GitHub 새 파일 링크는 승인된 publish 이후 사용 가능합니다.", ["paper/PAPER_TODO_2026-09-26.md", "AGENTS.md"])
    for i,(number,title,sub) in enumerate([("01","Verify references","Primary citations + author metadata"),("02","Format the manuscript","Venue template + figures + tables"),("03","Review the full diff","Approve before commit and publication")]):
        y=2.3+i*1.18
        text(s,.8,y,1,.62,number,33,BLUE,True)
        text(s,2.1,y,4.9,.48,title,24,INK,True)
        text(s,7.4,y+.06,5.1,.8,sub,18,MUTED)
        line(s,2.1,y+.85,12.45,y+.85)
    buf=BytesIO(); prs.save(buf)
    return deterministic_office(buf.getvalue()), slide_records


def build_docx():
    import cairosvg
    import markdown
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    doc=Document()
    metadata(doc.core_properties, PAPER.read_text().splitlines()[0].lstrip("# "))
    sec=doc.sections[0]
    sec.page_width, sec.page_height=Inches(8.27), Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.7)
    sec.left_margin=sec.right_margin=Inches(.7)
    sec.header_distance=sec.footer_distance=Inches(.32)
    for style in ("Normal", "Body Text", "List Bullet", "List Number"):
        st=doc.styles[style]; st.font.name="Liberation Serif"; st.font.size=Pt(11)
        st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.line_spacing=1.08
    for name,size in (("Title",23),("Heading 1",16),("Heading 2",12)):
        st=doc.styles[name]; st.font.name="Liberation Sans"; st.font.size=Pt(size)
        st.font.color.rgb=RGBColor.from_string(INK)
        st.paragraph_format.keep_with_next=True
    cap=doc.styles["Caption"]; cap.font.name="Liberation Serif"; cap.font.size=Pt(9)
    header=sec.header.paragraphs[0]
    header.text="MOTAR  ·  Research draft  ·  Evidence baseline 2da5b12"
    header.style="Caption"
    footer=sec.footer.paragraphs[0]; footer.text="26 September 2026  ·  Simulation-only evidence  |  "
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); footer._p.append(field)
    footer.style="Caption"
    tree=ET.fromstring("<article>"+markdown.markdown(PAPER.read_text(), extensions=["tables"])+"</article>")

    def add_inline(par, element, bold=False, italic=False):
        if element.text:
            run=par.add_run(element.text); run.bold=bold; run.italic=italic
        for child in element:
            if child.tag=="a":
                target=child.attrib["href"]
                if not target.startswith(("https://","http://","#")):
                    rel=(PAPER.parent/target).resolve().relative_to(ROOT).as_posix()
                    target="https://github.com/joshualikaist/MOTAR/blob/main/"+rel
                link=OxmlElement("w:hyperlink"); link.set(qn("r:id"),par.part.relate_to(target,RT.HYPERLINK,is_external=True))
                run=OxmlElement("w:r"); props=OxmlElement("w:rPr")
                color=OxmlElement("w:color"); color.set(qn("w:val"),BLUE); props.append(color); run.append(props)
                node=OxmlElement("w:t"); node.text="".join(child.itertext()); run.append(node); link.append(run); par._p.append(link)
            elif child.tag=="br": par.add_run().add_break()
            elif child.tag!="img": add_inline(par,child,bold or child.tag=="strong",italic or child.tag=="em")
            if child.tail: par.add_run(child.tail)

    for elem in tree:
        if elem.tag in ("h1","h2","h3"):
            level={"h1":0,"h2":1,"h3":2}[elem.tag]
            doc.add_heading("".join(elem.itertext()),level)
        elif elem.tag in ("ol","ul"):
            for item in elem:
                add_inline(doc.add_paragraph(style="List Number" if elem.tag=="ol" else "List Bullet"),item)
        elif elem.tag=="table":
            rows=elem.findall(".//tr")
            table=doc.add_table(rows=0, cols=len(list(rows[0])))
            table.style="Light Shading Accent 1"
            for i,row in enumerate(rows):
                cells=table.add_row().cells
                trpr=table.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if i==0: trpr.append(OxmlElement("w:tblHeader"))
                for cell,node in zip(cells,row):
                    p=cell.paragraphs[0]; add_inline(p,node,bold=i==0)
                    for run in p.runs: run.font.size=Pt(9)
                    p.paragraph_format.space_after=Pt(3)
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
        elif elem.tag=="p":
            img=elem.find("img")
            if img is not None:
                source=(PAPER.parent/img.attrib["src"]).resolve()
                data=cairosvg.svg2png(url=str(source),output_width=1800)
                p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
                pic=p.add_run().add_picture(BytesIO(data),width=Inches(6.65))
                pic._inline.docPr.set("descr",img.attrib.get("alt","MOTAR figure"))
            else:
                is_caption=elem.find("strong") is not None and (elem.find("strong").text or "").startswith("Figure ")
                add_inline(doc.add_paragraph(style="Caption" if is_caption else "Normal"),elem)
    buf=BytesIO(); doc.save(buf)
    return deterministic_office(buf.getvalue())


def build():
    v=evidence()
    pptx,slides=build_pptx(v)
    files={DECK:pptx,DOCX:build_docx()}
    source_paths=[PAPER.relative_to(ROOT).as_posix(),"presentations/build_companion.py","presentations/requirements-authoring.txt","docs/assets/paper/final/manifest.json"]
    source_paths += [p.relative_to(ROOT).as_posix() for p in sorted(FIGURES.glob("fig*.svg"))]
    for slide in slides:
        for path in slide["sources"]:
            if not (ROOT/path).is_file(): raise FileNotFoundError(path)
    record={"schema_version":1,"evidence_baseline":"2da5b12c1954ee768fcba59d1fb1fb9c78f7d536","date":STAMP,
            "generator":"presentations/build_companion.py","slide_count":len(slides),"slides":slides,
            "paper_word_count":len(PAPER.read_text().split()),
            "sources":{p:sha256((ROOT/p).read_bytes()).hexdigest() for p in source_paths},
            "files":{p:sha256(data).hexdigest() for p,data in files.items()}}
    files[MANIFEST]=(json.dumps(record,ensure_ascii=False,indent=2)+"\n").encode()
    return files


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check",action="store_true")
    args=parser.parse_args(); files=build()
    if args.check:
        problems=[p for p,data in files.items() if not (ROOT/p).is_file() or (ROOT/p).read_bytes()!=data]
        print(json.dumps({"status":"FAIL" if problems else "PASS","scope":"byte-exact DOCX, PPTX and manifest","problems":problems}))
        return bool(problems)
    for path,data in files.items():
        (ROOT/path).parent.mkdir(parents=True,exist_ok=True)
        (ROOT/path).write_bytes(data)
    print(json.dumps({"status":"BUILT","files":list(files)},indent=2))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
